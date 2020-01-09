import os
import csv
import openpyxl
from bs4 import BeautifulSoup

year = "2012"

input = "/media/SPE/processing/ua435/ua435_nDmescjNLkkjigbozTXSTY/masters/proceedings/" + str(year) +"/proceed/index.html"

f = open(input, "r", encoding='windows-1252', newline='')
soup = BeautifulSoup(f, 'html.parser')
pathRoot = "proceedings/" + str(year) + "/proceed/"
root = "/media/SPE/processing/ua435/ua435_nDmescjNLkkjigbozTXSTY/masters/"
sheet = []
wb = openpyxl.Workbook()

#make a sheet
worksheet = wb.active
#worksheet = wb.create_sheet(title=simpleTitle)
worksheet["H1"] = "Title"
worksheet["I1"] = str(year) + " Papers"
worksheet["H2"] = "Level"
worksheet["I2"] = "archival object"
worksheet["H3"] = "Ref ID"

#setup table headings
worksheet["A6"] = "ID"
worksheet["B6"] = "Location ID"
worksheet["C6"] = "Location"
worksheet["D6"] = "Container URI"
worksheet["E6"] = "Container"
worksheet["F6"] = "C#"
worksheet["G6"] = "Folder"
worksheet["H6"] = "F#"
worksheet["I6"] = "Title"
worksheet["J6"] = "Date 1 Display"
worksheet["K6"] = "Date 1 Normal"
worksheet["L6"] = "Date 2 Display"
worksheet["M6"] = "Date 2 Normal"
worksheet["N6"] = "Date 3 Display"
worksheet["O6"] = "Date 3 Normal"
worksheet["P6"] = "Date 4 Display"
worksheet["Q6"] = "Date 4 Normal"
worksheet["R6"] = "Date 5 Display"
worksheet["S6"] = "Date 5 Normal"
worksheet["T6"] = "Restrictions"
worksheet["U6"] = "General Note"
worksheet["V6"] = "Scope"
worksheet["W6"] = "DAO Filename"

#get table styles
tableStyle = openpyxl.worksheet.table.TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)

for para in soup.findAll('p'):
    if para.has_attr('style'):
        abstractIndex = None
        paperIndex = None
        supportingIndex = None
        
        #if para["style"] == "margin-left: 40; text-indent:-40; margin-bottom:-10":
        if "margin-left: .5in" in para["style"]:
            obj = []
            count = 0
            linkCount = 0
            check = False
            links = para.find_all('a')
            for text in para.stripped_strings:
                #print (text)
                count += 1
                if count == 1:
                    author = text
                if count == 2:
                    title = text
                    
            for link in links:
                if "abstract" in link.text.lower():
                    abstractIndex = linkCount
                if "paper" in link.text.lower():
                    paperIndex = linkCount
                if "supporting" in link.text.lower():
                    supportingIndex = linkCount
                linkCount += 1
                
            #print (author)
            #print ("\t" + title)
            #print (links)
            
            if not abstractIndex is None:
                #print (abstractIndex)
                abstract = links[abstractIndex]["href"]
                #print ("\t" + abstract)
            if not paperIndex is None:
                paper = links[paperIndex]["href"]
                #print ("\t" + paper)
            if not supportingIndex is None:
                supporting = links[supportingIndex]["href"]
                #print ("\t" + supporting)
            #print (count)
            obj.append(author + ", \"" + title + "\"")
            if not paperIndex is None:
                if not os.path.isfile(root + pathRoot + paper):
                    raise ValueError(paper)
                if not supportingIndex is None:
                    paper = paper + "|" + pathRoot + supporting
                    if not os.path.isfile(root + pathRoot + supporting):
                        raise ValueError(supporting)
                obj.append(pathRoot + paper)
            else:
                obj.append("")
            if not abstractIndex is None:
                abstractPath = root + pathRoot + abstract
                #print (abstractPath)
                #html = open(abstractPath, "r", encoding='windows-1252', newline='')
                #abstractText = html.read()
                #html.close()
                #print (abstractPath)
                try:
                    #print (abstractPath)
                    html = open(abstractPath, "r", newline='')
                    htmlSoup = BeautifulSoup(html, 'html.parser')
                    abstractText = htmlSoup.find("p").text.split("\n")[1]
                    html.close
                except:
                    #print (abstractPath)
                    html = open(abstractPath, "r", encoding='windows-1252', newline='')
                    htmlSoup = BeautifulSoup(html, 'html.parser')
                    abstractText = htmlSoup.find("p").text.split("\n")[1]
                    html.close
                abstractText = abstractText.strip('\n')
                abstractText = abstractText.strip('\r')
                #print (abstractText)
            
            
            if not abstractIndex is None:
                obj.append(abstractText)
                
            if not paperIndex is None:
                sheet.append(obj)
            
            if check == True:
                print (para.text)
            
            
                

f.close()

outputPath = '/media/SPE/processing/ua435/ua435_nDmescjNLkkjigbozTXSTY/metadata/' + str(year) + '.xlsx'

csvfile = open('/media/SPE/processing/ua435/ua435_nDmescjNLkkjigbozTXSTY/metadata/' + str(year) + "-papers.tsv", 'w', encoding='windows-1252', newline='')
writer = csv.writer(csvfile, delimiter='\t')
lineCount = 6
for line in sheet:
    lineCount += 1
    try:
        writer.writerow(line)
    except:
        print (line)
        writer.writerow([line[0], line[1]])
    worksheet["I" + str(lineCount)] = line[0]
    worksheet["W" + str(lineCount)] = line[1]
    
print ("Writing spreadsheet to " + outputPath)
				
table = openpyxl.worksheet.table.Table(ref='A6:W' + str(lineCount), displayName='Inventory', tableStyleInfo=tableStyle)
worksheet.add_table(table)

#styles for sheet info on top
worksheet["H1"].style = "Accent1"
worksheet["H2"].style = "Accent1"
worksheet["H3"].style = "Accent1"

#set column widths
worksheet.column_dimensions["I"].width = 60.0
worksheet.column_dimensions["F"].width = 15.0
worksheet.column_dimensions["J"].width = 15.0
worksheet.column_dimensions["K"].width = 15.0

wb.save(filename = os.path.join(outputPath))
    
csvfile.close()